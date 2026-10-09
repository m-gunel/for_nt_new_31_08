"""Парсер Excel: лист — единица, строка таблицы — карточка с шапкой, названием и листом.

Книга открывается один раз (read_only, data_only), каждый лист читается одним проходом.
Шапка — строки над первой строкой данных (до четырёх уровней); верхние уровни
протягиваются вправо: объединённые ячейки в режиме read_only не видны. Название таблицы —
строки из одной-двух ячеек над шапкой. Карточка строки:

    Таблица 12. Размещение филиалов… (лист «12»)
    Раздел: Уральский федеральный округ
    Федеральный округ: Челябинская область
    1 января 2025 — Всего: 613; ВСП КО…: 165

Даты — «1 января 2025»: так пишут запросы, и стеммер сопоставляет «январ» (ISO-дата даёт
токены «2025», «01», «01»). Числа — без хвостов float, по формату ячейки. Широкая строка
режется на несколько карточек ≤ BUDGET_TOKENS с повтором названия и подписи; сверху ещё
встанет «Документ: …». Листы не выбрасываются; лист без таблицы идёт строками текста.
Лимит — первые excel_config.max_rows + SCAN_ROWS строк листа, считая пустые: дальше
лист не читается. Легаси .xls — через Р7-Офис.
"""

import datetime as dt
import math
import os
import re
import shutil
import tempfile
from typing import Any, BinaryIO

import openpyxl
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException

from src.utils import get_logger
from src.utils.r7_office import R7Busy, convert_via_r7

from ._tables import Budget, pack_lines, split_parts
from .base import Parser, SkipFile, parser
from .doc_model import Block, ParsedDoc
from .metadata import ooxml_properties
from .ooxml_limits import XLSX_LARGE, XLSX_STREAMED, oversize_reason

logger = get_logger(__name__)

BUDGET_TOKENS = 450  # карточка; «Документ: <имя>» сверху уложит её в 500
SCAN_ROWS = 40  # где искать шапку
MAX_HEADER_ROWS = 4
PLACEHOLDERS = {"x", "х", "-", "–", "—", "…", "...", "*", "н/д"}
MONTHS = (
    "января", "февраля", "марта", "апреля", "мая", "июня",
    "июля", "августа", "сентября", "октября", "ноября", "декабря",
)  # fmt: skip
_SUP = str.maketrans("", "", "⁰¹²³⁴⁵⁶⁷⁸⁹")
# Сноска, приклеенная к слову: «Активы1», «капитал)2». Не трогаем «1к25», «9м24», «Н1», «МСФО 9».
_GLUED = re.compile(
    r"(?:(?<=[а-яё][а-яё])|(?<=\))|(?<=[а-яё],))\d{1,2}(?:,\d{1,2})*(?=[\s,.;:)]|$)"
)
_HYPHEN_NL = re.compile(r"(?<=[а-яёa-z])-\s*\n\s*(?=[а-яёa-z])")
_WS = re.compile(r"\s+")
_NUMBERING = re.compile(r"^\d{1,3}(\.\d{1,3})*\.?$")
_NOTE = re.compile(r"^\s*(\d{1,2}\s+\S|[¹²³⁴⁵⁶⁷⁸⁹*]|примечани|источник)", re.I)
_DEFAULT_SHEET = re.compile(r"^(лист|sheet)\s*\d*$", re.I)
_LITERAL = re.compile(r'"[^"]*"|\\.')  # текст в формате числа: «"…"» и «\x»


def _empty(v) -> bool:
    return v is None or (isinstance(v, str) and not v.strip())


def _number(v) -> bool:
    return isinstance(v, (int, float)) and not isinstance(v, bool)


def _yearlike(v) -> bool:
    return isinstance(v, int) and not isinstance(v, bool) and 1900 <= v <= 2100


def _placeholder(v) -> bool:
    return isinstance(v, str) and v.strip().lower() in PLACEHOLDERS


def clean_label(s) -> str:
    """Текст шапки, названия, подписи: переносы, сноски, пробелы."""
    s = _HYPHEN_NL.sub("", str(s)).translate(_SUP)
    return _WS.sub(" ", _GLUED.sub("", s)).strip()


def fmt_date(d) -> str:
    if isinstance(d, dt.time):
        return d.strftime("%H:%M")
    s = f"{d.day} {MONTHS[d.month - 1]} {d.year}"
    if isinstance(d, dt.datetime) and (d.hour or d.minute):
        s += f" {d:%H:%M}"
    return s


def fmt_number(v, number_format: str | None = None) -> str:
    """15.679387 → «15,68»; доля 0.0123 → «0,012»; формат «0.0%» → «12,5%»."""
    # «%» в кавычках или после «\» — просто знак: Excel на 100 не умножает.
    pct = "%" in _LITERAL.sub("", number_format or "")
    # Целое — тоже процент: ячейку 100 % (и 1.0) read_only отдаёт как int 1.
    if isinstance(v, int) and not pct:
        return str(v)
    if math.isnan(v) or math.isinf(v):
        return ""
    if pct:
        v *= 100
    a = abs(v)
    if a == 0:
        s = "0"
    elif a < 1e15 and v == round(v):
        s = str(int(round(v)))
    else:
        m = re.search(r"\.([0#]+)", (number_format or "").split(";")[0])
        dec = max(2, len(m.group(1)) if m else 0)
        if a < 1:  # малые доли — две значащие цифры, а не «0,00»
            dec = max(dec, -math.floor(math.log10(a)) + 1)
        s = f"{v:.{min(dec, 10)}f}".rstrip("0").rstrip(".")
    return s.replace(".", ",") + ("%" if pct else "")


def fmt_value(v, number_format: str | None = None) -> str:
    if isinstance(v, bool):
        return "да" if v else "нет"
    if _number(v):
        return fmt_number(v, number_format)
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return fmt_date(v)
    return _WS.sub(" ", str(v)).strip()


def header_text(v) -> str:
    if isinstance(v, (dt.datetime, dt.date, dt.time)):
        return fmt_date(v)
    if _number(v):
        return fmt_number(v)
    return clean_label(v)


# Значения, форматы, жирная одиночная, № строки.
Row = tuple[list[Any], list[str | None], bool, int]


def _read_rows(ws, limit: int) -> tuple[list[Row], bool]:
    """Строки листа без хвостовых пустых ячеек; второе — обрезан ли лист лимитом."""
    try:
        ws.reset_dimensions()  # иначе пустые ячейки со стилями раздувают лист до XFC
    except Exception:
        pass
    rows: list[Row] = []
    for r in ws.iter_rows():
        vals = [getattr(c, "value", None) for c in r]
        while vals and _empty(vals[-1]):
            vals.pop()
        fmts = [r[i].number_format if _number(v) else None for i, v in enumerate(vals)]
        filled = [i for i, v in enumerate(vals) if not _empty(v)]
        bold = False
        if len(filled) == 1:  # жирная одиночная строка — подпись раздела
            try:
                bold = bool(r[filled[0]].font.b)
            except Exception as exc:  # битый стиль ячейки — строка считается обычной
                logger.debug(
                    "xlsx: шрифт ячейки не прочитан (%s: %s)", type(exc).__name__, exc
                )
        number = next((c.row for c in r if getattr(c, "row", None)), None)
        rows.append((vals, fmts, bold, number or (rows[-1][3] + 1 if rows else 1)))
        if len(rows) >= limit:
            return rows, True
    return rows, False


def _nonempty(vals) -> int:
    return sum(1 for v in vals if not _empty(v))


def _is_data_row(vals) -> bool:
    ne = [v for v in vals if not _empty(v)]
    n_num = sum(1 for v in ne if _number(v) and not _yearlike(v))
    n_ph = sum(1 for v in ne if _placeholder(v))
    return len(ne) >= 2 and n_num >= 1 and (n_num + n_ph) * 3 >= len(ne) - 1


def _is_colnum_row(vals) -> bool:
    """Строка номеров граф «1 | 2 | 3 | 4» — не данные и не шапка."""
    ne = [v for v in vals if not _empty(v)]
    return (
        len(ne) >= 3
        and all(isinstance(v, int) for v in ne)
        and all(b - a == 1 for a, b in zip(ne, ne[1:]))
    )


def _titles_above(rows: list[Row], top: int, limit: int = 3) -> list[int]:
    """Название таблицы: до трёх ближайших непустых строк из 1–2 ячеек над шапкой."""
    out: list[int] = []
    for j in range(top - 1, -1, -1):
        n = _nonempty(rows[j][0])
        if n == 0:
            continue
        if n > 2 or len(out) >= limit:
            break
        out.insert(0, j)
    return out


def _layout(rows: list[Row]) -> tuple[list[int], int, list[int]] | None:
    """(строки шапки, начало данных, строки названия) или None — лист не таблица.

    Таблица начинается с первой строки из двух и более ячеек. Если ниже, в пределах
    MAX_HEADER_ROWS+2 строк, есть числовая строка данных, шапка — многоячеечные строки
    между ними. Иначе это текстовая таблица (шапка — первая строка, если следующие в
    тех же столбцах) или не таблица вовсе (оглавление, пояснения).
    """
    scan = rows[:SCAN_ROWS]
    top = next(
        (
            i
            for i, r in enumerate(scan)
            if _nonempty(r[0]) >= 2 and not _is_data_row(r[0])
        ),
        None,
    )
    first_data = next(
        (
            i
            for i, r in enumerate(scan)
            if _is_data_row(r[0]) and not _is_colnum_row(r[0])
        ),
        None,
    )
    if first_data is not None and (top is None or first_data < top):
        return [], first_data, _titles_above(rows, first_data)  # данные без шапки
    if top is None:
        return None
    if first_data is not None and first_data - top <= MAX_HEADER_ROWS + 2:
        header = [
            i
            for i in range(top, first_data)
            if _nonempty(rows[i][0]) >= 2 and not _is_colnum_row(rows[i][0])
        ][:MAX_HEADER_ROWS]
        return header, header[-1] + 1, _titles_above(rows, top)
    nxt = next(
        (
            i
            for i in range(top + 1, min(len(rows), top + 6))
            if _nonempty(rows[i][0]) >= 2
        ),
        None,
    )
    cols = {c for c, v in enumerate(rows[top][0]) if not _empty(v)}
    if (
        nxt is not None
        and len(cols & {c for c, v in enumerate(rows[nxt][0]) if not _empty(v)}) >= 2
    ):
        return [top], top + 1, _titles_above(rows, top)
    return None


def _headers(rows: list[Row], header_rows: list[int]) -> list[list[str]]:
    """Уровни шапки по столбцам сверху вниз; верхние уровни протянуты вправо."""
    width = max((len(rows[r][0]) for r in header_rows), default=0)
    grid = [
        [rows[r][0][c] if c < len(rows[r][0]) else None for c in range(width)]
        for r in header_rows
    ]
    for level in grid[:-1]:
        last = None
        for c in range(width):
            if not _empty(level[c]):
                last = level[c]
            elif last is not None:
                level[c] = last
    out = []
    for c in range(width):
        parts: list[str] = []
        for level in grid:
            if (
                not _empty(level[c])
                and (s := header_text(level[c]))
                and (not parts or parts[-1] != s)
            ):
                parts.append(s)
        out.append(parts)
    return out


def _label_columns(rows: list[Row], data_start: int, headers: list[str]) -> list[int]:
    """Столбцы подписи строки: номер п/п и первый текстовый столбец; следующий текстовый —
    тоже подпись, только если у него та же шапка («Показатель» над двумя столбцами) или её нет."""
    sample = [
        rows[i][0]
        for i in range(data_start, min(len(rows), data_start + 80))
        if _is_data_row(rows[i][0]) or _nonempty(rows[i][0]) >= 2
    ]
    width = max((len(r) for r in sample), default=0)
    labels: list[int] = []
    first_text_hdr = None
    for c in range(width):
        vals = [r[c] for r in sample if c < len(r) and not _empty(r[c])]
        if not vals:
            continue
        hdr = headers[c] if c < len(headers) else ""
        n_txt = sum(
            1
            for v in vals
            if isinstance(v, str)
            and not _placeholder(v)
            and not _NUMBERING.match(v.strip())
        )
        n_value = sum(
            1
            for v in vals
            if _placeholder(v)
            or isinstance(v, float)
            or (isinstance(v, int) and not 0 < v < 1000)
        )
        if n_value > 0.2 * len(vals):  # начались значения
            break
        if n_txt >= 0.6 * len(vals):
            if first_text_hdr is None:
                first_text_hdr = hdr
            elif hdr and hdr != first_text_hdr:
                break
            labels.append(c)
        elif first_text_hdr is None:  # номер п/п (1, 2, «3.1») перед текстом
            labels.append(c)
        else:
            break
    return labels


def _render_pairs(
    pairs: list[tuple[list[str], str]],
) -> list[tuple[str | None, list, str]]:
    """Пары с общим верхним уровнем шапки — одной строкой: «1 января 2025 — Всего: 613;
    ВСП КО…: 165». Одноуровневые — «шапка: значение»."""
    groups: list[list] = []
    for parts, val in pairs:
        top, rest = (
            (parts[0], " / ".join(parts[1:]))
            if len(parts) > 1
            else (None, " / ".join(parts))
        )
        if groups and groups[-1][0] == top:
            groups[-1][1].append((rest, val))
        else:
            groups.append([top, [(rest, val)]])
    out = []
    for top, items in groups:
        body = "; ".join(f"{r}: {v}" if r else v for r, v in items)
        out.append((top, items, f"{top} — {body}" if top else body))
    return out


def _unit_strings(pairs, prefix: str, budget: Budget) -> list[str]:
    """Группы пар, каждая вместе с prefix — в бюджете; широкая режется с повтором верхнего уровня."""
    units = []
    for top, items, text in _render_pairs(pairs):
        if budget.fits(prefix + text):
            units.append(text)
            continue
        lead = f"{top} — " if top else ""
        cur: list[str] = []
        for r, v in items:
            piece = f"{r}: {v}" if r else v
            if cur and not budget.fits(prefix + lead + "; ".join(cur + [piece])):
                units.append(lead + "; ".join(cur))
                cur = [piece]
            else:
                cur.append(piece)
        if cur:
            units.append(lead + "; ".join(cur))
    return units


def _sheet(ws, n: int, budget: Budget, limit: int) -> tuple[list[Block], str | None]:
    """Блоки листа и название его таблицы."""
    rows, cut = _read_rows(ws, limit + SCAN_ROWS)
    sheet = clean_label(ws.title)
    if cut:
        logger.warning(
            "лист «%s»: прочитаны первые %s строк (excel_config.max_rows %s + %s "
            "на шапку), остальные не индексируются",
            sheet,
            limit + SCAN_ROWS,
            limit,
            SCAN_ROWS,
        )
    if not rows:
        return [], None
    layout = _layout(rows)
    if layout is None:  # не таблица: строки как текст
        lines = [
            (clean_label(" | ".join(fmt_value(v) for v in r[0] if not _empty(v))), r[3])
            for r in rows
        ]
        blocks = pack_lines(
            f"Лист «{sheet}»\n",
            [x for x in lines if x[0]],
            budget,
            page=n,
            section=sheet,
        )
        return list(blocks), None
    header_rows, data_start, title_rows = layout
    hparts = _headers(rows, header_rows)
    headers = [" / ".join(p) for p in hparts]
    title = ". ".join(
        clean_label(" ".join(str(v) for v in rows[j][0] if not _empty(v)))
        for j in title_rows
    )
    title = re.sub(r"\.\.+", ".", title).strip()[:300]
    head = title or f"Лист «{sheet}»"
    if title and not _DEFAULT_SHEET.match(sheet) and sheet.lower() not in title.lower():
        head += f" (лист «{sheet}»)"
    labels = _label_columns(rows, data_start, headers)
    lab_hdr = next(
        (headers[c] for c in reversed(labels) if c < len(headers) and headers[c]), ""
    )
    place = {"page": n, "table_id": f"sheet{n}"}
    blocks: list[Block] = []
    section: list[str] = []
    notes: list[tuple[str, int]] = []
    last_was_data = False
    for vals, fmts, bold, number in rows[data_start:]:
        filled = _nonempty(vals)
        if filled == 0 or _is_colnum_row(vals):
            continue
        if filled == 1 and not any(_number(v) for v in vals):
            raw = next(v for v in vals if not _empty(v))
            text = clean_label(raw)
            if _NOTE.match(str(raw)) or len(text) > 250:
                notes.append((text, number))
            elif bold:
                section = [text] if last_was_data else (section + [text])[-2:]
                last_was_data = False
            # иначе подпись без значений (регион без КО) — пропускаем
            continue
        label = " ".join(
            header_text(vals[c])
            for c in labels
            if c < len(vals) and not _empty(vals[c])
        )
        pairs = [
            (hparts[c] if c < len(hparts) and hparts[c] else [f"столбец {get_column_letter(c + 1)}"],
             fmt_value(v, fmts[c]))
            for c, v in enumerate(vals)
            if c not in labels and not _empty(v) and not _placeholder(v)
        ]  # fmt: skip
        if not pairs and not label:
            continue
        sec = f"Раздел: {' / '.join(section)}\n"[:200] if section else ""
        row_line = f"{lab_hdr}: {label}" if lab_hdr and label else label
        if len(row_line) > 400:  # подпись-простыня не повторяется в каждом куске
            pairs.insert(0, ([], row_line))
            row_line = row_line[:200] + " …"
        last_was_data = True
        prefix = f"{head}\n{sec}" + (row_line + "\n" if row_line else "")
        path = " › ".join([sheet] + ([title] if title else []) + section)
        for text in split_parts(prefix, _unit_strings(pairs, prefix, budget), budget):
            blocks.append(
                Block(
                    text,
                    kind="table_row",
                    section=path,
                    row_start=number,
                    row_end=number,
                    whole=True,
                    **place,
                )
            )
    if section and not last_was_data:  # подписи без данных в конце — это примечания
        notes = [(t, rows[-1][3]) for t in section] + notes
    if notes:
        blocks += pack_lines(
            f"{head}\nПримечания:\n", notes, budget, page=n, section=sheet
        )
    return blocks, title or None


@parser("xlsx", "xls")
class XlsxParser(Parser):
    def parse_io(self, file: BinaryIO) -> ParsedDoc:
        # Формат узнаём по сигнатуре, а не по расширению: во вложениях почты имени нет.
        head = file.read(8)
        file.seek(0)
        if head.startswith(b"\xd0\xcf\x11\xe0"):
            return self._parse_legacy(file)
        if not head.startswith(b"PK"):
            # HTML или текст под видом .xls — не ошибка разбора, повторять незачем.
            logger.warning("файл не похож на книгу Excel (ни zip, ни OLE2) — пропущен")
            return ParsedDoc([])
        meta, titles = ooxml_properties(file, "xlsx")
        return self._parse_book(file, meta, titles)

    def _parse_legacy(self, file: BinaryIO) -> ParsedDoc:
        """Бинарный .xls openpyxl не читает — конвертируем штатным конвертером Р7-Офиса."""
        with tempfile.TemporaryDirectory() as tmp:
            src = os.path.join(tmp, "book.xls")
            dst = src + ".xlsx"
            with open(src, "wb") as out:
                shutil.copyfileobj(file, out)
            try:
                ok = convert_via_r7(src, dst, raise_busy=True)
            except R7Busy as exc:
                raise SkipFile(f"легаси .xls: Р7-Офис: {exc}", permanent=False) from exc
            if not ok:
                logger.warning(
                    "легаси .xls не конвертирован (Р7-Офис не найден или "
                    "не справился с файлом) — файл пропущен"
                )
                return ParsedDoc([])
            with open(dst, "rb") as converted:
                meta, titles = ooxml_properties(converted, "xlsx")
                return self._parse_book(converted, meta, titles)

    def _parse_book(self, file: BinaryIO, meta: dict, titles: list[str]) -> ParsedDoc:
        budget = Budget(BUDGET_TOKENS, self._config.ov_model.count_tokens)
        limit = self._config.excel_config.max_rows
        if reason := oversize_reason(file, XLSX_STREAMED, XLSX_LARGE):
            raise SkipFile(reason)
        try:
            wb = openpyxl.load_workbook(file, read_only=True, data_only=True)
        except (InvalidFileException, KeyError) as exc:  # zip, но не книга Excel
            raise SkipFile(f"не книга Excel ({type(exc).__name__}: {exc})") from exc
        blocks: list[Block] = []
        tables: list[str] = []
        sheets: list[str] = []
        try:
            for n, ws in enumerate(wb.worksheets, 1):
                if getattr(ws, "sheet_state", "visible") == "veryHidden":
                    continue
                sheets.append(clean_label(ws.title))
                try:
                    sheet_blocks, table = _sheet(ws, n, budget, limit)
                except (MemoryError, OSError):
                    raise  # временный сбой: книгу без листа нельзя отмечать готовой
                except Exception as exc:  # битый лист не должен ронять остальную книгу
                    logger.warning(
                        "лист «%s» не разобран: %s: %s",
                        ws.title,
                        type(exc).__name__,
                        exc,
                    )
                    continue
                blocks += sheet_blocks
                if table:
                    tables.append(table)
        finally:
            wb.close()
        # Название книги: свойства файла → название первой таблицы → имя листа (не «Лист1»).
        titles += tables[:1] + [s for s in sheets if not _DEFAULT_SHEET.match(s)][:1]
        return ParsedDoc(blocks, meta, titles, title_from_text=False)
