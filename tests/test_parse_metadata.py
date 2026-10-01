"""Метаданные документа и поля чанков на синтетических файлах всех форматов.

Файлы собираются в тесте (pymupdf, python-pptx, openpyxl, zipfile для docx) —
так видно, откуда берётся каждое поле. Чанки — через `file_chunks`, тот же путь,
что у индексатора.
"""

import hashlib
import os
import sys
import zipfile
from datetime import datetime, timedelta, timezone
from types import SimpleNamespace

import openpyxl
import pymupdf
import pytest
from pptx import Presentation

from src.parsers import file_chunks, read_document_text
from src.parsers import metadata as md


def _rows(path):
    rows = list(file_chunks(path))
    assert rows, "документ не дал чанков"
    return rows, rows[0]["meta"]


def _local(dt: datetime) -> datetime:
    return dt.astimezone().replace(tzinfo=None)


# ---------- PDF ----------


def _pdf(path, pages, toc=None, **metadata):
    with pymupdf.open() as doc:
        for text in pages:
            # Base-шрифты PyMuPDF латинские — текст тоже латиницей.
            doc.new_page().insert_text((72, 72), text)
        doc.set_metadata(metadata)
        if toc:
            doc.set_toc(toc)
        doc.save(str(path))


def test_pdf_metadata_pages_and_prefix(tmp_path):
    path = tmp_path / "report.pdf"
    _pdf(
        path,
        ["Quarterly risk report of the bank", "Second page about capital adequacy"],
        title="Quarterly Risk Report",
        author="Ivanov",
        subject="Risks",
        keywords="risk, capital",
        creationDate="D:20240115103000+03'00'",
        modDate="D:20240201120000Z",
    )
    rows, meta = _rows(path)
    assert meta["title"] == "Quarterly Risk Report"
    assert (meta["author"], meta["subject"], meta["keywords"]) == (
        "Ivanov",
        "Risks",
        "risk, capital",
    )
    assert meta["doc_created"] == _local(
        datetime(2024, 1, 15, 10, 30, tzinfo=timezone(timedelta(hours=3)))
    )
    assert meta["doc_modified"] == _local(
        datetime(2024, 2, 1, 12, 0, tzinfo=timezone.utc)
    )
    assert meta["page_count"] == 2
    assert meta["mime"] == "application/pdf"
    assert meta["sha256"] == hashlib.sha256(path.read_bytes()).hexdigest()
    assert meta["size"] == path.stat().st_size
    assert meta["language"] == "en"
    assert (
        meta["source_type"] == "file" and meta["source_uri"] == path.absolute().as_uri()
    )
    assert all(r["text"].startswith("Документ: report\n") for r in rows)
    assert {(r["page_start"], r["page_end"]) for r in rows} <= {(1, 1), (1, 2), (2, 2)}
    assert rows[0]["page_start"] == 1 and rows[-1]["page_end"] == 2


def test_pdf_junk_title_falls_back_to_bookmark(tmp_path):
    path = tmp_path / "scan.pdf"
    _pdf(
        path,
        ["Capital adequacy overview for the second quarter"],
        title="Microsoft Word - Document1.docx",
        toc=[[1, "Capital adequacy overview", 1]],
    )
    _, meta = _rows(path)
    assert meta["title"] == "Capital adequacy overview"


def test_pdf_title_not_confirmed_by_text_is_dropped(tmp_path):
    path = tmp_path / "notes.pdf"
    _pdf(path, ["Meeting notes about branch openings"], title="Annual Budget Plan")
    _, meta = _rows(path)
    assert meta["title"] == "Meeting notes about branch openings"  # первая строка


# ---------- DOCX ----------

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


def _docx(path, paragraphs, core, page_breaks=0):
    body = "".join(
        f"<w:p><w:r>{'<w:lastRenderedPageBreak/>' if i <= page_breaks and i else ''}<w:t>{p}</w:t></w:r></w:p>"
        for i, p in enumerate(paragraphs)
    )
    props = "".join(f"<{k}>{v}</{k}>" for k, v in core.items())
    with zipfile.ZipFile(path, "w") as z:
        z.writestr(
            "[Content_Types].xml",
            '<?xml version="1.0"?><Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types"/>',
        )
        z.writestr(
            "word/document.xml",
            f'<?xml version="1.0"?><w:document xmlns:w="{_W}"><w:body>{body}</w:body></w:document>',
        )
        z.writestr(
            "docProps/core.xml",
            '<?xml version="1.0"?><cp:coreProperties '
            'xmlns:cp="http://schemas.openxmlformats.org/package/2006/metadata/core-properties" '
            'xmlns:dc="http://purl.org/dc/elements/1.1/" xmlns:dcterms="http://purl.org/dc/terms/" '
            f'xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">{props}</cp:coreProperties>',
        )


def test_docx_core_properties_and_pages(tmp_path):
    path = tmp_path / "Положение.docx"
    _docx(
        path,
        [
            "Положение о закупках товаров и услуг",
            "Глава первая о порядке закупок",
            "Глава вторая о договорах",
        ],
        {
            "dc:title": "Положение о закупках",
            "dc:creator": "Иванова А. А.",
            "cp:lastModifiedBy": "Петров П. П.",
            "dc:subject": "Закупки",
            "dcterms:created": "2023-03-01T08:00:00Z",
            "dcterms:modified": "2023-04-02T09:30:00Z",
        },
        page_breaks=2,
    )
    rows, meta = _rows(path)
    assert meta["title"] == "Положение о закупках"
    assert (meta["author"], meta["last_modified_by"], meta["subject"]) == (
        "Иванова А. А.",
        "Петров П. П.",
        "Закупки",
    )
    assert meta["doc_created"] == _local(
        datetime(2023, 3, 1, 8, 0, tzinfo=timezone.utc)
    )
    assert meta["doc_modified"] == _local(
        datetime(2023, 4, 2, 9, 30, tzinfo=timezone.utc)
    )
    assert meta["page_count"] == 3  # два разрыва вёрстки Word → ≈3 страницы
    assert meta["mime"].endswith("wordprocessingml.document")
    assert meta["language"] == "ru"
    assert rows[0]["text"].startswith("Документ: Положение\n")


# ---------- PPTX ----------


def test_pptx_slides_and_title_from_first_slide(tmp_path):
    prs = Presentation()
    first = prs.slides.add_slide(prs.slide_layouts[0])
    first.shapes.title.text = "Стратегия развития"
    first.placeholders[1].text = "Подзаголовок презентации о стратегии банка"
    second = prs.slides.add_slide(prs.slide_layouts[1])
    second.shapes.title.text = "Цели на год"
    second.placeholders[
        1
    ].text = "Рост выручки и снижение затрат в регионах присутствия"
    prs.core_properties.title = "Презентация1"  # мусор — название берётся со слайда
    prs.core_properties.author = "Петров"
    prs.core_properties.last_modified_by = "Сидоров"
    path = tmp_path / "strategy.pptx"
    prs.save(str(path))
    rows, meta = _rows(path)
    assert meta["title"] == "Стратегия развития"
    assert (meta["author"], meta["last_modified_by"]) == ("Петров", "Сидоров")
    assert meta["page_count"] == 2
    assert rows[0]["page_start"] == 1 and rows[-1]["page_end"] == 2


# ---------- XLSX ----------


def test_xlsx_properties_and_sheet_count(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Отчёт"
    ws.append(["Показатель", "Значение"])
    for name, value in [
        ("Выручка компании", 100),
        ("Расходы компании", 70),
        ("Прибыль компании", 30),
    ]:
        ws.append([name, value])
    wb.create_sheet("Справка").append(["Примечание", "Данные за год"])
    wb.properties.title = "Выручка компании"
    wb.properties.creator = "Козлов"
    wb.properties.lastModifiedBy = "Орлов"
    path = tmp_path / "finance.xlsx"
    wb.save(str(path))
    rows, meta = _rows(path)
    assert meta["title"] == "Выручка компании"
    assert (meta["author"], meta["last_modified_by"]) == ("Козлов", "Орлов")
    assert meta["page_count"] == 2
    assert meta["mime"].endswith("spreadsheetml.sheet")
    # Лист с таблицей — карточки строк, лист «Справка» без чисел — текстом (не теряется).
    assert {r["kind"] for r in rows} == {"table_row", "text"}


# ---------- TXT ----------


def test_txt_lines_title_and_full_text(tmp_path):
    path = tmp_path / "check.txt"
    path.write_text(
        "Отчёт о проверке\r\n\r\n\r\nТекст первого абзаца отчёта о проверке филиала.\r\n"
        "Вторая строка того же абзаца с выводами.\r\n",
        encoding="utf-8",
    )
    rows, meta = _rows(path)
    assert meta["title"] == "Отчёт о проверке"
    assert (meta["encoding"], meta["language"], meta["mime"]) == (
        "utf-8",
        "ru",
        "text/plain",
    )
    assert rows[0]["line_start"] == 1 and rows[-1]["line_end"] == 5
    # /read_file: полный нормализованный текст, без префикса и перекрытий.
    assert read_document_text(path) == (
        "Отчёт о проверке\n\nТекст первого абзаца отчёта о проверке филиала.\n"
        "Вторая строка того же абзаца с выводами."
    )


# ---------- даты файловой системы ----------


def test_date_created_is_not_ctime(tmp_path):
    path = tmp_path / "old.txt"
    path.write_text("Старый документ с содержанием", encoding="utf-8")
    past = datetime(2020, 1, 1, 12, 0).timestamp()
    os.utime(path, (past, past))  # ctime становится «сейчас», а документ — из 2020 года
    meta = md.file_metadata(str(path))
    assert meta["date_modified"] == datetime(2020, 1, 1, 12, 0)
    if meta["date_created"] is not None:
        assert meta["date_created"] <= meta["date_modified"]
        assert meta["date_created"] != datetime.fromtimestamp(path.stat().st_ctime)


def test_unknown_birth_time_gives_empty_date(tmp_path, monkeypatch):
    path = tmp_path / "a.txt"
    path.write_text("текст", encoding="utf-8")
    monkeypatch.setattr(sys, "platform", "sunos5")
    assert md.birth_time(str(path), SimpleNamespace(st_mtime=0)) is None
    monkeypatch.setattr(md, "birth_time", lambda *a: None)
    assert md.file_metadata(str(path))["date_created"] is None


@pytest.mark.skipif(not sys.platform.startswith("linux"), reason="statx — только Linux")
def test_statx_birth_time_on_linux(tmp_path):
    path = tmp_path / "new.txt"
    path.write_text("x", encoding="utf-8")
    born = md._statx_btime(str(path))
    if born is not None:  # ФС может не хранить время рождения (tmpfs старых ядер)
        assert abs((datetime.now() - born).total_seconds()) < 60


# ---------- поля по отдельности ----------


@pytest.mark.parametrize(
    ("title", "head", "expected"),
    [
        (
            "Microsoft Word - Отчёт о рисках.docx",
            "годовой отчёт о рисках банка",
            "Отчёт о рисках",
        ),
        ("ПОЛОЖЕНИЕ О ЗАЧЁТЕ", "положение о зачете требований", "ПОЛОЖЕНИЕ О ЗАЧЁТЕ"),
        ("Документ1", "любой текст", None),
        ("PLEASE READ FIRST", "please read first", None),
        ("Untitled", "untitled", None),
        ("C:\\Users\\ivanov\\отчёт.doc", "отчёт", None),
        ("/home/user/отчёт.pdf", "отчёт", None),
        ("Бюджет развития на год", "совсем другой текст о погоде", None),
        ("", "текст", None),
    ],
)
def test_clean_title(title, head, expected):
    assert md.clean_title(title, head) == expected


def test_pdf_and_ooxml_dates():
    plus3 = timezone(timedelta(hours=3))
    assert md.parse_pdf_date("D:20240115103000+03'00'") == _local(
        datetime(2024, 1, 15, 10, 30, tzinfo=plus3)
    )
    assert md.parse_pdf_date("D:2024") == datetime(2024, 1, 1)
    assert md.parse_pdf_date("D:19691231") is None  # неправдоподобно — пусто
    assert md.parse_pdf_date("мусор") is None
    assert md.parse_w3cdtf("2023-03-01T08:00:00Z") == _local(
        datetime(2023, 3, 1, 8, tzinfo=timezone.utc)
    )
    assert md.parse_w3cdtf("2023-03-01") == datetime(2023, 3, 1)
    assert md.parse_w3cdtf("01.03.2023") is None


def test_language_and_mime():
    assert md.language("Отчёт о деятельности банка за год") == "ru"
    assert md.language("Annual report of the bank for the year") == "en"
    assert md.language("Отчёт банка annual report за 2024 year") == "ru,en"
    assert md.language("12 34") is None
    assert md.sniff_mime(b"%PDF-1.7", "pdf") == "application/pdf"
    assert md.sniff_mime(b"PK\x03\x04", "xlsx").endswith("spreadsheetml.sheet")
    assert md.sniff_mime(b"\xd0\xcf\x11\xe0", "doc") == "application/msword"
    assert md.sniff_mime(b"plain", "txt") == "text/plain"
    assert md.sniff_mime(b"\x00\x01", "bin") is None
