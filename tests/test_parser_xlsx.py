"""Парсер xlsx (src/parsers/xlsx.py) на синтетических книгах openpyxl.

Строка таблицы — карточка: название таблицы, подпись строки и пары «шапка: значение»
с датами словами. Листы не выбрасываются, широкая строка режется с повтором названия.
"""

import datetime as dt
import logging
from pathlib import Path

import openpyxl
import pytest
from openpyxl.styles import Font

from src.parsers import file_chunks
from src.parsers.xlsx import XlsxParser, fmt_number, fmt_value

LEGACY_XLS = Path(__file__).parent / "fixtures" / "legacy.xls"
TITLE = "Таблица 12. Размещение филиалов и внутренних структурных подразделений (ВСП)"
VSP = "ВСП КО, зарегистрированных в данном регионе, единиц"


def _book(path: Path, sheets: dict[str, list[list]]) -> Path:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for row in rows:
            ws.append(row)
    wb.save(path)
    return path


def _texts(path: Path) -> list[str]:
    return [c["text"] for c in XlsxParser().iter_chunks(str(path))]


def _regions_table() -> list[list]:
    return [
        [None, TITLE],
        [],
        [
            "Федеральный округ",
            dt.datetime(2025, 1, 1),
            None,
            dt.datetime(2026, 1, 1),
            None,
        ],
        [None, "Всего", VSP, "Всего", VSP],
        ["Челябинская область", 613, 165, 576, 162],
        ["Свердловская область", 900, 200, 880, 190],
    ]


def test_two_level_header_gives_readable_card(tmp_path):
    texts = _texts(_book(tmp_path / "obs.xlsx", {"12": _regions_table()}))
    card = next(t for t in texts if "Челябинская" in t)
    assert card == (
        f"{TITLE}\n"
        "Федеральный округ: Челябинская область\n"
        f"1 января 2025 — Всего: 613; {VSP}: 165\n"
        f"1 января 2026 — Всего: 576; {VSP}: 162"
    )
    joined = "\n".join(texts)
    assert (
        "Column_" not in joined
        and "2025-01-01" not in joined
        and "00:00:00" not in joined
    )


def test_title_and_sheet_in_every_card_and_fields(tmp_path):
    path = _book(tmp_path / "book.xlsx", {"Отчёт": _regions_table()})
    rows = list(file_chunks(path))
    cards = [r for r in rows if r["kind"] == "table_row"]
    assert len(cards) == 2
    for r in cards:
        assert r["text"].startswith(f"Документ: book\n{TITLE} (лист «Отчёт»)\n")
        assert (r["page_start"], r["table_id"]) == (1, "sheet1")
        assert r["section"] == f"Отчёт › {TITLE}"
    assert [r["row_start"] for r in cards] == [5, 6]  # номера строк Excel
    assert rows[0]["meta"]["title"] == TITLE  # свойств нет — название первой таблицы
    assert rows[0]["meta"]["page_count"] == 1


def test_numeric_sheet_is_not_skipped(tmp_path):
    rows = [["Год", "Показатель A", "Показатель B"]] + [
        [2000 + i, i * 1.5, i * 2.25] for i in range(1, 60)
    ]
    texts = _texts(_book(tmp_path / "nums.xlsx", {"Данные": rows}))
    assert len(texts) >= 50  # прежний парсер пропускал такой лист по доле текста


def test_bold_single_row_is_a_section(tmp_path):
    path = tmp_path / "sections.xlsx"
    _book(path, {"4": [
        [None, "Таблица 4. Структура активов"],
        ["№ п.п.", "Показатель (млрд руб.)", dt.datetime(2026, 2, 1)],
        [None, "Активы"],
        ["1", "Денежные средства", 12913.797468],
        [None, "Пассивы"],
        ["2", "Средства клиентов", 9050.5],
    ]})  # fmt: skip
    wb = openpyxl.load_workbook(path)
    for row in (3, 5):
        wb["4"].cell(row=row, column=2).font = Font(bold=True)
    wb.save(path)
    rows = list(file_chunks(path))
    money = next(r for r in rows if "Денежные средства" in r["text"])
    clients = next(r for r in rows if "Средства клиентов" in r["text"])
    assert (
        "Раздел: Активы\n" in money["text"]
        and "1 февраля 2026: 12913,8" in money["text"]
    )
    assert money["section"] == "4 › Таблица 4. Структура активов › Активы"
    assert "Раздел: Пассивы\n" in clients["text"]


def test_value_formatting():
    assert fmt_number(15.679387) == "15,68"
    assert fmt_number(-95.588501) == "-95,59"
    assert fmt_number(0.0123) == "0,012"
    assert fmt_number(2.0) == "2"
    assert fmt_number(0.125, "0.0%") == "12,5%"
    assert fmt_value(dt.datetime(2025, 1, 1)) == "1 января 2025"
    assert fmt_value(dt.datetime(2025, 3, 8, 14, 30)) == "8 марта 2025 14:30"


def test_wide_row_is_split_with_title_and_label(tmp_path, config):
    header = ["Показатель"] + [
        f"Длинное название столбца номер {i}" for i in range(200)
    ]
    data = ["Строка один"] + [i * 1.5 + 0.25 for i in range(200)]
    texts = _texts(_book(tmp_path / "wide.xlsx", {"Широкий": [header, data]}))
    assert len(texts) > 1
    for t in texts:
        assert "Показатель: Строка один" in t
        assert config.ov_model.count_tokens(t) <= 450
    joined = "\n".join(texts)
    assert "Длинное название столбца номер 199: 298,75" in joined


def test_row_limit_cuts_sheet_with_warning(tmp_path, config, monkeypatch, caplog):
    monkeypatch.setattr(config.excel_config, "max_rows", 100)
    rows = [["Номер", "Сумма"]] + [[f"Счёт {i}", i + 0.5] for i in range(300)]
    with caplog.at_level(logging.WARNING):
        texts = _texts(_book(tmp_path / "long.xlsx", {"Реестр": rows}))
    assert 50 < len(texts) < 300  # лист не потерян, но обрезан
    assert any("больше 100 строк" in r.message for r in caplog.records)


def test_sheet_without_table_is_text(tmp_path):
    texts = _texts(
        _book(
            tmp_path / "readme.xlsx",
            {
                "Пояснения": [
                    ["Пояснения к отчётности банка"],
                    ["Данные приведены в млрд рублей"],
                ]
            },
        )
    )
    assert texts == [
        "Лист «Пояснения»\nПояснения к отчётности банка\nДанные приведены в млрд рублей"
    ]


def test_empty_book_and_html_under_xls_name(tmp_path, caplog):
    assert _texts(_book(tmp_path / "empty.xlsx", {"Лист1": []})) == []
    fake = tmp_path / "report.xls"
    fake.write_text("<html><table><tr><td>1</td></tr></table></html>", encoding="utf-8")
    with caplog.at_level(logging.WARNING):
        assert _texts(fake) == []  # не исключение: повторять разбор незачем
    assert any("ни zip, ни OLE2" in r.message for r in caplog.records)


@pytest.mark.skipif(not LEGACY_XLS.exists(), reason="нет фикстуры legacy.xls")
def test_legacy_xls_goes_through_r7_converter(tmp_path, monkeypatch):
    converted = _book(tmp_path / "converted.xlsx", {"12": _regions_table()})
    stub = tmp_path / "x2t"
    stub.write_text(f'#!/bin/sh\ncp "{converted}" "$2"\n', encoding="utf-8")
    stub.chmod(0o755)
    monkeypatch.setenv("LQRAG_R7_X2T", str(stub))
    texts = _texts(LEGACY_XLS)
    assert any("Челябинская область" in t and "1 января 2025" in t for t in texts)
