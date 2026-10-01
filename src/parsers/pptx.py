from collections.abc import Generator
from typing import Any, BinaryIO
from zipfile import BadZipFile

from pptx import Presentation

from src.utils import get_logger

from .base import Chunk, Parser, parser

logger = get_logger(__name__)


@parser("pptx")
class PptxParser(Parser):
    def iter_chunks_io(self, file: BinaryIO) -> Generator[Chunk, Any, None]:
        idx = 0
        texts = []
        try:
            prs = Presentation(file)
        except BadZipFile:
            logger.warning(
                "Failed to read PPTX file, it may not be a valid PPTX format."
            )
            return
        for slide in prs.slides:
            for shape in slide.shapes:
                if txt := getattr(shape, "text", None):
                    texts.append(txt)
        for split in self._config.embed_splitter.split_text("\n".join(texts)):
            yield Chunk(idx=idx, text=split)
            idx += 1
