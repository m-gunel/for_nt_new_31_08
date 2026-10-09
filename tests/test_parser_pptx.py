"""Разбор pptx (src/parsers/pptx.py) на презентациях, собранных python-pptx в тесте."""

from pathlib import Path

from pptx import Presentation
from pptx.enum.shapes import PP_PLACEHOLDER
from pptx.oxml.ns import qn
from pptx.util import Inches

from src.parsers import file_chunks


def _box(shapes, text: str, top: int = 1):
    box = shapes.add_textbox(Inches(1), Inches(top), Inches(6), Inches(1))
    box.text_frame.text = text
    return box


def _deck(path: Path) -> Path:
    prs = Presentation()
    blank = prs.slide_layouts[6]
    s1 = prs.slides.add_slide(prs.slide_layouts[5])  # только заголовок
    s1.shapes.title.text = "Кредиты для бизнеса"
    group = s1.shapes.add_group_shape()
    _box(group.shapes, "Ставка от 12 процентов годовых для малого бизнеса", 2)
    _box(group.shapes, "Срок кредита до пяти лет без залога", 3)
    s1.notes_slide.notes_text_frame.text = "Подчеркнуть отсутствие залога при выдаче"
    s2 = prs.slides.add_slide(blank)
    _box(s2.shapes, "Тарифы по отделениям")
    table = s2.shapes.add_table(3, 3, Inches(1), Inches(2), Inches(6), Inches(2)).table
    for r, row in enumerate(
        [
            ["Отделение", "Код", "Телефон"],
            ["Москва, Тверская 1", "ZX01", "495-111"],
            ["Казань, Баумана 5", "ZX02", "843-222"],
        ]
    ):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    s3 = prs.slides.add_slide(blank)
    steps = _box(s3.shapes, "Подать заявку онлайн").text_frame
    steps.add_paragraph().text = "Получить решение за день"
    steps.add_paragraph().text = "Подписать договор в отделении"
    for para in steps.paragraphs:
        para._p.get_or_add_pPr().append(
            para._p.makeelement(qn("a:buAutoNum"), {"type": "arabicPeriod"})
        )
    prs.save(str(path))
    return path


def test_groups_tables_notes_and_slide_numbers(tmp_path):
    rows = list(file_chunks(_deck(tmp_path / "кредиты.pptx")))
    joined = "\n".join(r["text"] for r in rows)
    # Текст фигур внутри группы не теряется.
    assert "Ставка от 12 процентов годовых для малого бизнеса" in joined
    assert "Срок кредита до пяти лет без залога" in joined
    notes = [r for r in rows if r["kind"] == "notes"]
    assert [n["page_start"] for n in notes] == [1]
    assert notes[0]["text"].endswith(
        "Кредиты для бизнеса\nЗаметки: Подчеркнуть отсутствие залога при выдаче"
    )
    tables = [r for r in rows if r["kind"] == "table"]
    assert (
        len(tables) == 1
        and tables[0]["page_start"] == 2
        and tables[0]["table_id"] == "s2t1"
    )
    assert tables[0]["text"].endswith(
        "Тарифы по отделениям\nОтделение | Код | Телефон\n"
        "Москва, Тверская 1 | ZX01 | 495-111\nКазань, Баумана 5 | ZX02 | 843-222"
    )
    # Заголовок слайда без фигуры-заголовка — короткая первая строка.
    assert tables[0]["section"] == "Тарифы по отделениям"
    assert rows[0]["section"] == "Кредиты для бизнеса"
    assert rows[0]["meta"]["title"] == "Кредиты для бизнеса"
    assert rows[0]["meta"]["page_count"] == 3


def test_auto_numbering_is_restored(tmp_path):
    rows = list(file_chunks(_deck(tmp_path / "шаги.pptx")))
    slide3 = "\n".join(r["text"] for r in rows if r.get("page_start") == 3)
    assert (
        "1. Подать заявку онлайн\n2. Получить решение за день\n3. Подписать договор в отделении"
        in slide3
    )


def test_typical_deck_chunks_are_unchanged(tmp_path):
    """Обычные слайды: правки заголовков, номеров и шапок их текст не трогают."""
    rows = [
        (r["kind"], r["page_start"], r["section"], r["text"])
        for r in file_chunks(_deck(tmp_path / "кредиты.pptx"))
        if r["kind"] != "card"
    ]
    assert rows == [
        (
            "text",
            1,
            "Кредиты для бизнеса",
            "Документ: кредиты\nКредиты для бизнеса\n"
            "Ставка от 12 процентов годовых для малого бизнеса\n"
            "Срок кредита до пяти лет без залога",
        ),
        (
            "notes",
            1,
            "Кредиты для бизнеса",
            "Документ: кредиты\nКредиты для бизнеса\n"
            "Заметки: Подчеркнуть отсутствие залога при выдаче",
        ),
        (
            "table",
            2,
            "Тарифы по отделениям",
            "Документ: кредиты\nТарифы по отделениям\nОтделение | Код | Телефон\n"
            "Москва, Тверская 1 | ZX01 | 495-111\nКазань, Баумана 5 | ZX02 | 843-222",
        ),
        (
            "text",
            3,
            "Тарифы по отделениям",
            "Документ: кредиты\n1. Подать заявку онлайн\n2. Получить решение за день\n"
            "3. Подписать договор в отделении",
        ),
    ]


def _table_slide(prs, title: str, rows: list[list[str]]):
    slide = prs.slides.add_slide(prs.slide_layouts[5])  # только заголовок
    slide.shapes.title.text = title
    table = slide.shapes.add_table(
        len(rows), len(rows[0]), Inches(1), Inches(2), Inches(6), Inches(2)
    ).table
    for r, row in enumerate(rows):
        for c, value in enumerate(row):
            table.cell(r, c).text = value
    return slide


def test_divider_slide_between_tables_is_indexed(tmp_path):
    prs = Presentation()
    _table_slide(
        prs,
        "Объём кредитного портфеля",
        [["Сегмент", "2025", "2026"], ["Корпоративный", "120", "150"]],
    )
    divider = prs.slides.add_slide(prs.slide_layouts[0])
    divider.shapes.title.text = "Раздел 2. Риски кредитного портфеля"
    _table_slide(prs, "Динамика просрочки",
                 [["Сегмент", "2025", "2026"], ["Корпоративный", "1,2%", "1,5%"]])  # fmt: skip
    path = tmp_path / "риски.pptx"
    prs.save(str(path))
    rows = [r for r in file_chunks(path) if r["kind"] != "card"]
    divider_rows = [r for r in rows if "Риски кредитного портфеля" in r["text"]]
    assert [r["page_start"] for r in divider_rows] == [2]
    # Заголовок слайда с таблицей — первая строка её чанка, отдельным чанком не идёт.
    assert [r["kind"] for r in rows if "Динамика просрочки" in r["text"]] == ["table"]
    table = next(r for r in rows if r.get("page_start") == 3)
    assert table["text"] == (
        "Документ: риски\nДинамика просрочки\nСегмент | 2025 | 2026\n"
        "Корпоративный | 1,2% | 1,5%"
    )


def test_kpi_numbers_are_kept_and_slide_number_is_not(tmp_path):
    prs = Presentation()
    layout = prs.slide_layouts[5]
    slide = prs.slides.add_slide(layout)
    slide.shapes.title.text = "Выручка компании"
    for k, text in enumerate(["1500", "млн рублей за год рост", "85%"], 2):
        _box(slide.shapes, text, k)
    number = next(
        ph for ph in layout.placeholders
        if ph.placeholder_format.type == PP_PLACEHOLDER.SLIDE_NUMBER
    )  # fmt: skip
    slide.shapes.clone_placeholder(number)
    slide.placeholders[number.placeholder_format.idx].text_frame.text = "7"
    path = tmp_path / "итоги.pptx"
    prs.save(str(path))
    [row] = [r for r in file_chunks(path) if r["kind"] != "card"]
    assert row["text"] == (
        "Документ: итоги\nВыручка компании\n1500\nмлн рублей за год рост\n85%"
    )


def test_long_first_row_of_slide_table(tmp_path, config):
    count = config.ov_model.count_tokens
    long_cell = (
        "Настоящий регламент определяет порядок взаимодействия подразделений банка "
        "при рассмотрении заявок клиентов на получение кредитных продуктов. "
    ) * 32
    prs = Presentation()
    _table_slide(
        prs,
        "Паспорт процесса",
        [["Цель", long_cell]]
        + [
            [f"Шаг {i}", f"Исполнитель {i} передаёт документы за {i} дня"]
            for i in range(1, 16)
        ],
    )
    path = tmp_path / "паспорт.pptx"
    prs.save(str(path))
    tables = [r["text"] for r in file_chunks(path) if r["kind"] == "table"]
    assert count(long_cell) > 450
    assert len(tables) <= 6  # без правки — сотни чанков «шапка + слово»
    assert all(count(t) <= 500 for t in tables)  # 450 + строка «Документ: …»
    joined = "\n".join(tables)
    assert all(f"Шаг {i} | Исполнитель {i}" in joined for i in range(1, 16))
